from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from docx import Document

root = Path(r"C:\Users\bonkc\Documents\wound-splat")
original = root / "MAIN DOCUMENT (WOUND-SPLAT) - WORKING COPY.docx"
revised = root / "MAIN DOCUMENT (WOUND-SPLAT) - REVISED.docx"
change_log = root / "MAIN DOCUMENT (WOUND-SPLAT) - CHANGE LOG.docx"


def check_zip(path):
    with ZipFile(path) as zf:
        bad = zf.testzip()
        assert bad is None, f"Corrupt ZIP member in {path.name}: {bad}"
        return {
            name: sha256(zf.read(name)).hexdigest()
            for name in zf.namelist()
            if name.startswith("word/media/")
        }


original_media = check_zip(original)
revised_media = check_zip(revised)
check_zip(change_log)
assert original_media == revised_media, "Embedded media changed or was lost"

doc = Document(revised)
log_doc = Document(change_log)
body_text = "\n".join(p.text for p in doc.paragraphs)

assert len(doc.paragraphs) == 543
assert len(doc.tables) == 23
assert len(doc.sections) == 1
assert len(doc.inline_shapes) == 14
assert "Algorithm Implementation" in body_text
assert "Step 7 – Persistence, Visualization, and Reporting" in body_text
assert "This patient-code lookup is not secure authentication." in body_text
assert "Longitudinal registration and Wagner and University of Texas classification support remain planned Thesis 2 extensions." in body_text
assert "T_{process}" not in body_text
assert "\\frac" not in body_text
assert "Table X" not in body_text
assert "TODO" not in body_text
assert "TBD" not in body_text
assert "to be confirmed" not in body_text.lower()

native_equations = doc.element.xpath(".//m:oMathPara")
assert len(native_equations) == 15, f"Expected 15 native equations, found {len(native_equations)}"

formula_labels = [p for p in doc.paragraphs if p.text.strip().startswith("Formula ")]
assert len(formula_labels) == 15
for p in formula_labels:
    assert p.alignment == 0, f"Formula label is not left aligned: {p.text}"
    assert all(r.font.name == "Times New Roman" and r.font.size.pt == 12 and r.bold for r in p.runs)

formula_titles = {
    "Processing Time", "Registration Rate", "Percentage Error", "Mean Absolute Percentage Error",
    "Scale Calibration", "Dice Similarity Coefficient", "Intersection over Union",
    "Pipeline Completion Rate", "Relative Difference for Benefit Criteria", "Benefit-Criterion Rank",
    "Relative Difference for Cost Criteria", "Cost-Criterion Rank", "Criterion Score",
    "Weighted Criterion Score", "Total Weighted Score",
}
title_paragraphs = [p for p in doc.paragraphs if p.text.strip() in formula_titles]
assert len(title_paragraphs) == 15
for p in title_paragraphs:
    assert p.alignment == 0
    assert all(r.font.name == "Times New Roman" and r.font.size.pt == 12 and r.italic for r in p.runs)

algorithm = []
collecting = False
for p in doc.paragraphs:
    if p.text.strip() == "Algorithm Implementation":
        collecting = True
    if collecting and p.text.strip() == "Multiple Constraints and Trade-off Analysis":
        break
    if collecting:
        algorithm.append(p)
assert len(algorithm) == 16
for p in algorithm:
    assert p.paragraph_format.line_spacing == 2.0
    for r in p.runs:
        assert r.font.name == "Times New Roman" and r.font.size.pt == 12

for index in range(12, 22):
    table = doc.tables[index]
    assert len(table.columns) == 6, f"Survey table {index + 1} does not have 6 columns"
    scale = [cell.text.strip() for cell in table.rows[1].cells[1:]]
    assert scale == ["5", "4", "3", "2", "1"], f"Survey table {index + 1} scale is {scale}"

log_text = "\n".join(p.text for p in log_doc.paragraphs)
assert "96 logged changes" in log_text, "Change-log count is missing or incorrect"

print("DOCX package integrity: PASS")
print(f"Embedded media preserved: PASS ({len(revised_media)} files; {len(doc.inline_shapes)} inline shapes)")
print(f"Structure: PASS ({len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(doc.sections)} section)")
print("Questionnaire scales: PASS (10 tables, 5-point scale)")
print("Algorithm Implementation and Thesis 1/2 boundaries: PASS")
print("Native editable equations: PASS (15 equations)")
print("Formula labels and italic titles: PASS")
print("Algorithm Implementation formatting: PASS (Times New Roman 12 pt, double spaced)")
print("Copyable plain-text formula residue check: PASS")
print("Placeholder/obsolete wording check: PASS")
print("Change log package: PASS")
