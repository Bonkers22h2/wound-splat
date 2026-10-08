from collections import Counter
from pathlib import Path
import sys
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"C:\Users\bonkc\Documents\wound-splat\MAIN DOCUMENT (WOUND-SPLAT) - WORKING COPY.docx")
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(r"C:\Users\bonkc\Documents\wound-splat\docx_inventory.txt")


def iter_blocks(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


doc = Document(SOURCE)
lines = [
    f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)} inline_shapes={len(doc.inline_shapes)}",
    f"styles={Counter(p.style.name for p in doc.paragraphs).most_common(30)}",
]

pidx = 0
tidx = 0
for bidx, block in enumerate(iter_blocks(doc), start=1):
    if isinstance(block, Paragraph):
        pidx += 1
        text = block.text.replace("\t", " <TAB> ").replace("\n", " ").strip()
        if text:
            lines.append(f"B{bidx:04d} P{pidx:04d} [{block.style.name}] {text}")
    else:
        tidx += 1
        lines.append(f"B{bidx:04d} TABLE{tidx:03d} rows={len(block.rows)} cols={len(block.columns)}")
        for ridx, row in enumerate(block.rows, start=1):
            cells = [" / ".join(p.text.strip() for p in c.paragraphs if p.text.strip()).replace("\t", " <TAB> ") for c in row.cells]
            lines.append(f"  R{ridx:02d}: " + " || ".join(cells))

OUT.write_text("\n".join(lines), encoding="utf-8")
print("\n".join(lines[:5]))
print(f"wrote {OUT} with {len(lines)} lines")
