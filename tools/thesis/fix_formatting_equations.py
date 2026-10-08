from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(r"C:\Users\bonkc\Documents\wound-splat")
DOCX = ROOT / "MAIN DOCUMENT (WOUND-SPLAT) - REVISED.docx"
LOG = ROOT / "MAIN DOCUMENT (WOUND-SPLAT) - CHANGE LOG.docx"

doc = Document(DOCX)


def paragraphs_between(start_text, end_text):
    collecting = False
    result = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text == start_text:
            collecting = True
        if collecting:
            if text == end_text:
                break
            result.append(p)
    return result


def set_run_font(run, *, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_text_format(p, *, bold=False, italic=False, centered=False, first_indent=False, keep=False):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Inches(0.5) if first_indent else Inches(0)
    pf.keep_with_next = keep
    for run in p.runs:
        set_run_font(run, bold=bold, italic=italic)


# Match the added Algorithm Implementation section to the surrounding thesis text.
algorithm_paragraphs = paragraphs_between("Algorithm Implementation", "Multiple Constraints and Trade-off Analysis")
for p in algorithm_paragraphs:
    text = p.text.strip()
    if text == "Algorithm Implementation":
        p.style = doc.styles["Heading 2"]
        set_text_format(p, bold=True, keep=True)
    elif text.startswith("Step "):
        p.style = doc.styles["Heading 3"]
        set_text_format(p, bold=True, keep=True)
    else:
        p.style = doc.styles["normal"]
        set_text_format(p, first_indent=True)


def m_el(tag):
    return OxmlElement(f"m:{tag}")


def m_run(text):
    r = m_el("r")
    wrpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), "Times New Roman")
    wrpr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    wrpr.append(size)
    r.append(wrpr)
    t = m_el("t")
    t.text = text
    if text.startswith(" ") or text.endswith(" "):
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    return r


def as_nodes(value):
    if isinstance(value, str):
        return [m_run(value)]
    if isinstance(value, list):
        nodes = []
        for item in value:
            nodes.extend(as_nodes(item))
        return nodes
    return [value]


def sub(base, subscript):
    node = m_el("sSub")
    e = m_el("e")
    for child in as_nodes(base):
        e.append(child)
    s = m_el("sub")
    for child in as_nodes(subscript):
        s.append(child)
    node.extend([e, s])
    return node


def frac(numerator, denominator):
    node = m_el("f")
    num = m_el("num")
    den = m_el("den")
    for child in as_nodes(numerator):
        num.append(child)
    for child in as_nodes(denominator):
        den.append(child)
    node.extend([num, den])
    return node


def accent_hat(base):
    node = m_el("acc")
    prop = m_el("accPr")
    char = m_el("chr")
    char.set(qn("m:val"), "̂")
    prop.append(char)
    e = m_el("e")
    for child in as_nodes(base):
        e.append(child)
    node.extend([prop, e])
    return node


def summation(lower, upper, expression):
    node = m_el("nary")
    prop = m_el("naryPr")
    char = m_el("chr")
    char.set(qn("m:val"), "∑")
    loc = m_el("limLoc")
    loc.set(qn("m:val"), "undOvr")
    prop.extend([char, loc])
    low = m_el("sub")
    high = m_el("sup")
    body = m_el("e")
    for child in as_nodes(lower):
        low.append(child)
    for child in as_nodes(upper):
        high.append(child)
    for child in as_nodes(expression):
        body.append(child)
    node.extend([prop, low, high, body])
    return node


xhat_i = sub(accent_hat("x"), "i")
xref_i = sub("x", "ref,i")

FORMULAS = [
    ("Formula 1: Processing Time", "Processing Time", [sub("T", "process"), " = ", sub("t", "completed"), " − ", sub("t", "started")]),
    ("Formula 2: Registration Rate", "Registration Rate", [sub("R", "reg"), " = ", frac(sub("N", "registered"), sub("N", "extracted")), " × 100"]),
    ("Formula 3: Percentage Error", "Percentage Error", [sub("E", "pct"), " = ", frac(["|", accent_hat("x"), " − ", sub("x", "ref"), "|"], ["|", sub("x", "ref"), "|"]), " × 100"]),
    ("Formula 4: Mean Absolute Percentage Error", "Mean Absolute Percentage Error", ["MAPE = ", frac("1", "n"), " × ", summation("i = 1", "n", ["|", frac([deepcopy(xhat_i), " − ", deepcopy(xref_i)], deepcopy(xref_i)), "|"]), " × 100"]),
    ("Formula 5: Scale Calibration", "Scale Calibration", ["s = ", frac(sub("L", "ref,real"), sub("L", "ref,model"))]),
    ("Formula 6: Dice Similarity Coefficient", "Dice Similarity Coefficient", ["Dice = ", frac("2TP", "2TP + FP + FN")]),
    ("Formula 7: Intersection over Union", "Intersection over Union", ["IoU = ", frac("TP", "TP + FP + FN")]),
    ("Formula 8: Pipeline Completion Rate", "Pipeline Completion Rate", [sub("R", "complete"), " = ", frac(sub("N", "completed"), sub("N", "submitted")), " × 100"]),
    ("Formula 9: Relative Difference for Benefit Criteria", "Relative Difference for Benefit Criteria", [sub("RD", "benefit"), " = ", frac("B − V", "B")]),
    ("Formula 10: Benefit-Criterion Rank", "Benefit-Criterion Rank", [sub("R", "benefit"), " = 10 × (1 − ", sub("RD", "benefit"), ")"]),
    ("Formula 11: Relative Difference for Cost Criteria", "Relative Difference for Cost Criteria", [sub("RD", "cost"), " = ", frac("V − B", "V")]),
    ("Formula 12: Cost-Criterion Rank", "Cost-Criterion Rank", [sub("R", "cost"), " = 10 × (1 − ", sub("RD", "cost"), ")"]),
    ("Formula 13: Criterion Score", "Criterion Score", ["CS = ", frac("R", "10")]),
    ("Formula 14: Weighted Criterion Score", "Weighted Criterion Score", ["WCS = CS × CW"]),
    ("Formula 15: Total Weighted Score", "Total Weighted Score", [sub("S", "total"), " = ", summation("j = 1", "m", ["(", sub("CS", "j"), " × ", sub("CW", "j"), ")"])]),
]


def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def add_native_equation(p, elements):
    clear_paragraph(p)
    p.style = doc.styles["normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Inches(0)
    math_para = m_el("oMathPara")
    props = m_el("oMathParaPr")
    just = m_el("jc")
    just.set(qn("m:val"), "center")
    props.append(just)
    math_para.append(props)
    math = m_el("oMath")
    for element in as_nodes(elements):
        math.append(element)
    math_para.append(math)
    p._p.append(math_para)


def insert_after(p, text=""):
    new_p = OxmlElement("w:p")
    p._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    result = Paragraph(new_p, p._parent)
    if text:
        result.add_run(text)
    return result


for old_label, title, elements in FORMULAS:
    labels = [p for p in doc.paragraphs if p.text.strip() == old_label]
    if len(labels) != 1:
        raise RuntimeError(f"Expected one label {old_label!r}, found {len(labels)}")
    label = labels[0]
    equation = label._p.getprevious()
    if equation is None or equation.tag != qn("w:p"):
        raise RuntimeError(f"Equation paragraph missing before {old_label}")

    number = old_label.split(":", 1)[0]
    clear_paragraph(label)
    label.add_run(number)
    label.style = doc.styles["normal"]
    set_text_format(label, bold=True, keep=True)

    title_p = insert_after(label, title)
    title_p.style = doc.styles["normal"]
    set_text_format(title_p, italic=True, keep=True)

    equation.getparent().remove(equation)
    title_p._p.addnext(equation)
    from docx.text.paragraph import Paragraph
    equation_p = Paragraph(equation, title_p._parent)
    add_native_equation(equation_p, elements)

doc.save(DOCX)


# Extend the existing detailed change log with the formatting corrections.
log_doc = Document(LOG)
table = log_doc.tables[0]
entries = [
    (
        "Chapter 3 – Algorithm Implementation",
        "Normalized the added section to Times New Roman, 12-point text with double spacing, APA-style paragraph indentation, and the existing heading hierarchy.",
        "The added content previously inherited default Word formatting instead of the thesis document's direct formatting.",
    ),
    (
        "Chapter 3 – Formulas 1–15",
        "Separated each formula number and italicized title, then converted every displayed expression into a centered, editable native Word equation with structured fractions, subscripts, accents, and summation limits.",
        "This follows the document's formula presentation pattern and prevents LaTeX or plain-text equation residue when opened in Word or Google Docs.",
    ),
]
start = len(table.rows)
for offset, (location, change, reason) in enumerate(entries, start=0):
    row = table.add_row()
    values = [str(start + offset), location, change, reason]
    for cell, value in zip(row.cells, values):
        cell.text = value
        for p in cell.paragraphs:
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                set_run_font(run)

note = log_doc.add_paragraph("Formatting follow-up: The revised document now contains 96 logged changes, including the Algorithm Implementation style correction and conversion of all 15 formulas to native equations.")
for run in note.runs:
    set_run_font(run)
log_doc.save(LOG)

print(f"updated {DOCX}")
print(f"updated {LOG}")
print("algorithm paragraphs formatted:", len(algorithm_paragraphs))
print("native equations created:", len(FORMULAS))
