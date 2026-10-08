from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SRC = Path(r"C:\Users\bonkc\Documents\wound-splat\MAIN DOCUMENT (WOUND-SPLAT) - WORKING COPY.docx")
OUT = Path(r"C:\Users\bonkc\Documents\wound-splat\MAIN DOCUMENT (WOUND-SPLAT) - REVISED.docx")
LOG_OUT = Path(r"C:\Users\bonkc\Documents\wound-splat\MAIN DOCUMENT (WOUND-SPLAT) - CHANGE LOG.docx")

doc = Document(SRC)
changes = []


def log(location, old, new, reason):
    changes.append({"location": location, "old": old, "new": new, "reason": reason})


def find_paragraph(exact):
    matches = [p for p in doc.paragraphs if p.text.strip() == exact.strip()]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph for {exact[:80]!r}, found {len(matches)}")
    return matches[0]


def set_paragraph_text(p, text):
    # These edited paragraphs do not contain drawings. Keep the first run's
    # character formatting and the paragraph's style/spacing/indentation.
    if any("<w:drawing" in r._element.xml for r in p.runs):
        raise RuntimeError(f"Refusing to erase a drawing in paragraph: {p.text[:80]}")
    if p.runs:
        first = p.runs[0]
        first.text = text
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)


def replace(location, old, new, reason):
    p = find_paragraph(old)
    set_paragraph_text(p, new)
    log(location, old, new, reason)
    return p


def insert_before(reference, text, style, location, reason):
    p = reference.insert_paragraph_before(text)
    p.style = style
    log(location, "[Missing]", text, reason)
    return p


def set_cell_text(cell, text):
    p = cell.paragraphs[0]
    set_paragraph_text(p, text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def remove_last_table_column(table):
    # The questionnaire header merges the six rating cells. Decrease that
    # grid-span instead of deleting the entire merged header cell.
    for tr in table._tbl.tr_lst:
        last_tc = tr.tc_lst[-1]
        span = last_tc.tcPr.gridSpan
        if span is not None and int(span.val) > 1:
            span.val = int(span.val) - 1
        else:
            tr.remove(last_tc)
    grid = table._tbl.tblGrid
    grid.remove(grid.gridCol_lst[-1])


# Chapter 1: correct only claims that conflict with the implemented scope.
replace(
    "Chapter 1 – Background of the Study",
    "Diabetes is a major public health concern in the Philippines, with approximately 1 in 14 Filipino adults diagnosed with the condition (International Diabetes Federation, 2024). Among its most debilitating complications is the diabetic foot ulcer (DFU), which, if not properly monitored, often leads to lower-limb amputation. Globally, DFUs affect approximately 18.6 million people each year and precede 80% of lower extremity amputations in diabetic patients. The five-year mortality rate for a patient with a DFU is approximately 30%, rising to over 70% for those who undergo major amputation.",
    "Diabetes is a major public health concern in the Philippines, with approximately 1 in 14 Filipino adults diagnosed with the condition (International Diabetes Federation, 2024). Among its most debilitating complications is the diabetic foot ulcer (DFU), which, if not properly monitored, may lead to lower-limb amputation. Globally, DFUs affect approximately 18.6 million people each year and precede many lower-extremity amputations in people with diabetes. The five-year mortality rate after a DFU is approximately 30% and is higher after major amputation (Armstrong et al., 2023).",
    "Added the existing Armstrong reference to the global-burden claims and removed wording that implied an inevitable outcome.",
)

replace(
    "Chapter 1 – Background of the Study",
    "Image segmentation is a critical step in automated DFU analysis, serving as the bridge between raw 2D images and clinically meaningful 3D measurements. Accurate Segmentation is essential for isolating the wound boundary and classifying tissue types such as granulation and fibrin, granulation, and callus, which provide key indicators of healing or deterioration. Deep learning has significantly advanced wound segmentation, with architectures such as HarDNet-MSEG combined with the Segment Anything Model (SAM) achieving an average Dice of 83.44% on the DFUC 2022 dataset (Chen et al., 2023). Public datasets and challenges such as DFUC 2022, DFUC 2023, and DFUC 2024 have driven progress by providing standardized benchmarks for algorithm comparison.  While most research focuses on 2D segmentation, recent work on WoundNeRF has explored 3D-consistent segmentation from multi-view images using Neural Radiance Fields (Chierchia et al., 2026). For the Wound-Splat framework, the wound region is isolated from the reconstructed point cloud, while a trained network characterizes tissue type on selected frames.",
    "Image segmentation is a critical step in automated DFU analysis, serving as the bridge between raw 2D images and clinically meaningful information. Accurate segmentation is essential for isolating the wound boundary and classifying tissue types such as fibrin, granulation, and callus, which provide indicators of healing or deterioration. Deep learning has significantly advanced wound segmentation, with architectures such as HarDNet-MSEG combined with the Segment Anything Model (SAM) achieving an average Dice of 83.44% on the DFUC 2022 dataset (Chen et al., 2023). Public datasets and challenges such as DFUC 2022, DFUC 2023, and DFUC 2024 have driven progress by providing standardized benchmarks for algorithm comparison. While most research focuses on 2D segmentation, recent work on WoundNeRF has explored 3D-consistent segmentation from multi-view images using Neural Radiance Fields (Chierchia et al., 2026). In Wound-Splat, the wound region is isolated from the reconstructed point cloud, while a separate trained network characterizes tissue type within an operator-defined region of a selected frame.",
    "Removed the duplicated tissue class, corrected capitalization, and separated the 3D isolation and 2D tissue-analysis functions.",
)

replace(
    "Chapter 1 – Background of the Study",
    "To address these issues, the proposed Wound-Splat framework aims to provide a comprehensive method for 3D volumetric analysis of diabetic foot ulcers using only standard smartphone videos and existing GPU-accelerated 3D Gaussian Splatting implementations. Unlike existing approaches, the framework integrates 3D reconstruction with 2D wound segmentation for tissue characterization, supports longitudinal monitoring through 3D model registration across visits, and maps measurements to established DFU classification systems (Wagner, University of Texas) for clinically interpretable outputs. While the system requires a GPU-equipped workstation for reconstruction, this is still a fraction of the cost of a professional medical 3D scanner and can be centrally deployed in a hospital IT department rather than purchased per clinic. The application handles video upload and result viewing, making the system accessible without requiring specialized training.",
    "To address these issues, the proposed Wound-Splat framework investigates 3D volumetric analysis of diabetic foot ulcers using standard smartphone videos and an existing GPU-accelerated 3D Gaussian Splatting implementation. The implemented prototype integrates 3D reconstruction and geometric measurement with separate 2D tissue characterization. Longitudinal registration across visits and evidence-based Wagner and University of Texas classification support remain planned Thesis 2 components. The current system requires a compatible NVIDIA GPU backend, while the web application handles video upload, progress monitoring, result viewing, and report access. Users still require capture and operating instructions to obtain suitable inputs and interpret the outputs correctly.",
    "Separated implemented functions from planned Thesis 2 functions and removed unsupported claims about cost and training requirements.",
)

replace(
    "Chapter 1 – Objective 1",
    "Implement a 3D reconstruction pipeline utilizing Structure from Motion (SfM) and stock 3D Gaussian Splatting (3DGS) to accurately extract volumetric DFU models from smartphone videos.",
    "Implement a 3D reconstruction pipeline utilizing Structure from Motion (SfM) and stock 3D Gaussian Splatting (3DGS) to produce measurable wound-like models from smartphone videos and evaluate its geometric accuracy under defined test conditions.",
    "Replaced the unverified accuracy claim with an explicit evaluation requirement.",
)
replace(
    "Chapter 1 – Objective 2",
    "Design a temporal tracking module utilizing geometric registration algorithms to align 3D models across visits, automating the calculation of wound volume changes over time.",
    "Design and, in Thesis 2, implement a temporal tracking module using geometric registration to align 3D models across visits and calculate changes in wound measurements over time.",
    "Marked longitudinal tracking as a planned Thesis 2 component.",
)
replace(
    "Chapter 1 – Objective 3",
    "Automate the extraction of spatial metrics (depth, volume, area) to programmatically map 3D reconstruction outputs to established Wagner and University of Texas classification systems.",
    "Automate the extraction of spatial metrics, including planar wound opening area, cavity volume, and maximum depth, and, in Thesis 2, develop classification support that combines these outputs with the additional clinician-observed information required by the Wagner and University of Texas systems.",
    "Prevented geometric measurements from being presented as sufficient for complete clinical classification.",
)
replace(
    "Chapter 1 – Objective 4",
    "Develop and evaluate a client-server web application using Next.js and FastAPI for GPU-delegated 3DGS processing, validating algorithmic accuracy with simulated models and assessing overall software quality utilizing the ISO/IEC 25010 Model.",
    "Develop a client-server web application using Next.js and FastAPI for GPU-backed 3DGS processing; verify the geometric measurement calculations with synthetic point-cloud phantoms; and prepare the ISO/IEC 25010:2023 software-quality evaluation for Thesis 2.",
    "Clarified that synthetic tests verify the measurement routine only and that the full software-quality evaluation is scheduled for Thesis 2.",
)

replace(
    "Chapter 1 – Significance of the Study",
    "Furthermore, the integration of 3D reconstruction measurements with established DFU classification systems (Wagner, University of Texas) represents a novel contribution. While these classification systems are widely used in clinical practice, they have traditionally relied on manual measurement and visual assessment. By mapping automated 3D measurements to these systems, this research explores how computer vision can bridge the gap between technical reconstruction and clinical utility.",
    "The planned integration of 3D reconstruction measurements with the Wagner and University of Texas systems may support more structured documentation. However, complete classification also depends on clinical findings such as infection, ischemia, and anatomical involvement that cannot be inferred from geometric measurements alone. The Thesis 1 contribution therefore centers on the implemented reconstruction, measurement, tissue-characterization, visualization, and reporting workflow; classification support remains a Thesis 2 objective.",
    "Reframed the classification contribution so it does not claim an unfinished or clinically incomplete feature.",
)
replace(
    "Chapter 1 – Significance of the Study",
    "The primary beneficiaries of the research would be diabetic patients, who could directly upload videos through a mobile or web interface and receive 3D wound assessments without requiring specialized clinic visits; Healthcare professionals, who could obtain objective 3D measurements (volume, depth, area) to support clinical decision-making; Hospitals and specialty clinics, which could perform volumetric wound analysis without the need for dedicated medical 3D scanners; and researchers in computer vision and medical imaging, who could build upon the framework and findings presented in this study. The methodology for 3D wound reconstruction, longitudinal monitoring, and classification integration may inform future work in wound assessment and related medical imaging applications.",
    "Potential beneficiaries include patients who may use the web portal to submit capture data under appropriate clinical guidance; healthcare professionals who may review reconstructed models, geometric measurements, tissue-characterization outputs, and reports as decision-support information; healthcare institutions investigating lower-cost wound-documentation workflows; and researchers studying smartphone-based reconstruction and wound assessment. The prototype does not replace clinical examination, and its use does not eliminate the need for clinic visits, trained operation, or clinician judgment.",
    "Removed unsupported promises about mobile access, clinic avoidance, and operation without training.",
)

replace(
    "Chapter 1 – Scope",
    "The scope encompasses the entire pipeline from video capture to clinical output – a Structure from Motion (SfM) pipeline for camera pose estimation and sparse point cloud generation, followed by 3DGS-based dense reconstruction, image wound segmentation, comparison of 3D models from multiple patient visits to document wound change over time, and the mapping of all measured parameters to a rule-based classification based on established DFU classification systems (Wagner, University of Texas) to produce clinically interpretable outputs. An absolute scale is obtained from a reference object of known size placed beside the wound during capture. The system detects a standard bank card or coin in the video frames and derives the number of centimeters represented by one unit of the reconstructed model. Where no reference object is detected, measurements are reported as uncalibrated and flagged as approximate rather than presented as clinical values. The framework’s output is a clinically useful report containing measurements, classifications, and change-over-time visualizations to support clinical decision-making. Moreover, validation will be based on simulated wound models with verifiable baseline measurements. Moreover, validation will be based on synthetic wound phantoms with analytically known dimensions so that measured values can be compared against exact ground truth. This follows the principle of phantom-based evaluation used in the SALVE benchmark (Chierchia et al., 2025a), which employed physical rather than synthetic phantoms. Lastly, the framework will be implemented on a web or mobile application for video upload and result viewing, following a client-server architecture model where reconstruction and segmentation processing occurs on a backend GPU workstation. The framework is delivered as a web application with two user roles. A patient portal allows a patient to sign in with an assigned code, upload a wound video, follow processing status, and view the resulting model and report. A clinical administration view allows staff to register patients and monitor the processing queue. Completed scans can be inspected in an interactive three-dimensional viewer in the browser, and the measurements are compiled into a generated PDF report.",
    "The implemented scope covers a responsive client-server web application that accepts a smartphone video and processes it through frame extraction, Structure from Motion camera estimation and sparse reconstruction, stock 3D Gaussian Splatting optimization, preview rendering, three-dimensional wound-region isolation, scale calibration, geometric measurement, and PDF report generation. The browser supports video upload, progress monitoring, result viewing, and operator interaction, while the GPU-equipped backend performs reconstruction and model inference. A supported bank card or coin placed beside the captured subject provides the physical scale; when reliable calibration is unavailable, the scan may still complete but its size-dependent outputs are identified as uncalibrated. Measurements comprise planar wound opening area, cavity volume, maximum depth, width, and height. A separate two-dimensional tissue module analyzes an operator-defined rectangular region and reports fibrin, granulation, and callus percentages when sufficient wound pixels are detected. Synthetic point-cloud phantoms verify the geometric measurement calculations only. Controlled physical end-to-end validation, respondent-based software-quality evaluation, longitudinal registration, and Wagner and University of Texas classification support will be completed in Thesis 2.",
    "Removed duplicated sentences and separated the implemented Thesis 1 scope from planned Thesis 2 work.",
)
replace(
    "Chapter 1 – Scope",
    "Processing runs as a fixed seven-stage pipeline: frame extraction, camera-pose estimation, 3DGS optimization, preview rendering, wound isolation, scale calibration and measurement, and report generation. Each stage reports progress so that a scan in process can be monitored rather than appearing unresponsive. A complete scan takes approximately thirty to sixty minutes on a single consumer GPU.",
    "Processing runs as a fixed seven-stage pipeline: frame extraction, camera-pose estimation, 3DGS optimization, preview rendering, wound isolation, scale calibration and measurement, and report generation. Each stage reports progress so that a scan in process can be monitored rather than appearing unresponsive. Processing time varies with video length, registered-frame count, reconstruction density, hardware, and system load and will be measured during the controlled Thesis 2 evaluation.",
    "Removed the unsupported fixed 30–60 minute runtime claim.",
)
replace(
    "Chapter 1 – Delimitation",
    "The framework is subject to several delimitations. The study utilizes existing open-source libraries and models for implementing the algorithms used within the framework; the contribution lies in the application, integration, and evaluation of these latest methods for DFU monitoring. Validation is limited to phantom-based evaluation using simulated wound models and geometric objects; clinical validation with real patient data is outside the study scope. Reconstruction is evaluated using smartphone-captured video under controlled lighting conditions, and the framework is not validated for uncontrolled environments or varying lighting. The system is designed for a single NVIDIA RTX series GPU, and cloud or distributed computing is not included. The system does not include integration with electronic medical records to existing clinics. Classification suggestions are based on identified tissue types present in the wound and geometric measurements – the framework does not assess infection, ischemia, or other clinical factors that are part of complete classification systems. Factors such as skin temperature and blood flow, which are important clinical indicators, cannot be captured through standard RGB video and are not measured by the framework.",
    "The framework uses existing open-source reconstruction and learning components; the study contribution lies in their adaptation, integration, and evaluation for wound-assessment support. No prospective patient wound videos or clinical validation are included. Smartphone-video reconstruction is limited to controlled capture conditions and has not been validated for uncontrolled lighting, motion, glare, exudate, or occlusion. The primary execution environment is a single NVIDIA RTX 4050 Laptop GPU with 6 GB of video memory. A compatible cloud GPU may be provisioned manually for high-memory runs, but the application does not provide automatic failover, distributed execution, or automatic job offloading. Electronic medical-record integration, secure authentication, and role-based access control are outside the current prototype. Synthetic point-cloud tests verify the measurement mathematics but do not validate the complete capture-to-measurement pipeline. Longitudinal registration and Wagner and University of Texas classification support remain planned; the current RGB workflow does not assess infection, ischemia, perfusion, skin temperature, or other clinical findings required for complete classification.",
    "Resolved the cloud-computing contradiction and stated the actual validation, security, and clinical limitations.",
)

# Chapter 2: remove editorial wording and fix scope/citations/standards.
replace(
    "Chapter 2 – Conceptual Framework",
    "The conceptual framework presents Wound-Splat as an Input–Process–Output system. The model is revised from the placeholder IPO figure in the submitted document so that the relationship between the research inputs, computational stages, and clinical outputs is explicit. The framework is not intended to replace clinical examination. Instead, it positions the proposed software as a measurement and decision-support tool whose outputs must be interpreted together with clinical findings.",
    "The conceptual framework presents Wound-Splat as an Input–Process–Output system in which research inputs are transformed through defined computational stages into measurable and reviewable outputs. The framework is not intended to replace clinical examination. It positions the proposed software as a measurement and decision-support tool whose outputs must be interpreted together with clinical findings.",
    "Removed an internal editorial note about revising a placeholder figure.",
)
replace(
    "Chapter 2 – Conceptual Framework",
    "The input consists primarily of a short smartphone video of the wound and the information needed to associate the recording with a particular visit. A scale reference is important because reconstruction from uncalibrated images may recover shape without an absolute physical scale. The process begins with frame extraction and camera-pose estimation, followed by 3D reconstruction and wound-region segmentation. The reconstructed representation is filtered and converted into a measurable surface. The output consists of geometric measurements, visualizations, tissue-related information, longitudinal comparisons, and classification suggestions. The final report should state the measurement conditions and should clearly distinguish automated results from clinical diagnoses.",
    "The input consists primarily of a smartphone video, the information needed to associate the recording with a scan, a selected scale-reference type, and an operator-defined tissue-analysis region when tissue characterization is requested. The implemented process performs frame extraction, camera-pose estimation, 3DGS optimization, wound-region isolation, scale calibration, heightfield-based geometric measurement, visualization, tissue inference, and report generation. Current outputs include the reconstructed representation, previews, calibration status, planar wound opening area, cavity volume, maximum depth, width, height, tissue percentages when sufficient wound pixels are detected, and a PDF report. Longitudinal comparisons and classification support are planned Thesis 2 outputs and are not presented as current results.",
    "Aligned the IPO description with the implemented pipeline and identified future outputs explicitly.",
)
replace(
    "Chapter 2 – Research Paradigm",
    "The revised research paradigm follows the Input–Process–Output model described above. The input is the smartphone-captured video and associated wound information. The process is the sequence of image preparation, SfM, segmentation, 3DGS optimization, surface reconstruction, geometric analysis, registration, classification mapping, and software delivery. The output is a clinically interpretable report that contains a reconstructed wound model, measurements, tissue characterization, visual previews, and change-over-time information. This arrangement directly connects the proposed system to the objectives of the study.",
    "The research paradigm follows the Input–Process–Output model described above. The implemented process covers image preparation, SfM, 3DGS optimization, three-dimensional wound isolation, scale calibration, geometric analysis, operator-guided two-dimensional tissue characterization, visualization, persistence, and report generation. Component-level verification examines the measurement calculations using synthetic point clouds and evaluates tissue segmentation using labeled DFUTissue images. Controlled physical end-to-end evaluation and respondent-based software-quality evaluation will be performed in Thesis 2. Longitudinal registration and classification support will also be implemented and evaluated in Thesis 2.",
    "Removed editorial wording and separated implemented processing, current verification, and planned evaluation.",
)
replace(
    "Chapter 2 – Structure from Motion",
    "Structure from Motion estimates camera poses and sparse scene geometry from multiple images. The method depends on feature detection, feature matching, geometric verification, triangulation, and bundle adjustment. Recent implementations of incremental SfM systems, such as those utilized in modern open-source pipelines like COLMAP, continue to prioritize robustness, accuracy, completeness, and scalability in complex environments (Liu et al., 2024). This source is important for Wound-Splat because COLMAP is used to estimate camera poses and generate a sparse point cloud before densification or Gaussian optimization.",
    "Structure from Motion estimates camera poses and sparse scene geometry from multiple images. The method depends on feature detection, feature matching, geometric verification, triangulation, and bundle adjustment. COLMAP implements a general-purpose incremental SfM and Multi-View Stereo pipeline designed for robust image-based reconstruction (Schönberger & Frahm, 2016). This source is important for Wound-Splat because COLMAP is used to estimate camera poses and generate the sparse initialization required before Gaussian optimization.",
    "Replaced a mismatched citation with the foundational COLMAP reference used by the implementation.",
)
replace(
    "Chapter 2 – Wound Segmentation",
    "Encoder–decoder networks, heavily popularized by foundational architectures like U-Net and extensively reviewed in recent biomedical segmentation literature (Z. Huang et al., 2025), remain highly effective when annotated medical images are limited. Wound-splat may use encoder-decoder networks, popularized by U-Net and reviewed in recent biomedical segmentation literature (Huang et al., 2025b), and remain effective when annotated medical images are limited. General-purpose models such as HarDNet-MSEG, SAM, and SAM 2 were considered but not adopted, because tissue characterization requires a model trained on labeled wound tissue rather than one assumed to generalize from unrelated biomedical images.",
    "Encoder–decoder networks, including U-Net-style architectures reviewed in recent biomedical segmentation literature (Z. Huang et al., 2025), remain effective when annotated medical images are limited. Wound-Splat uses a U-Net-style model with a MiT-b3 encoder and scSE attention for tissue characterization. General-purpose approaches such as HarDNet-MSEG, SAM, and SAM 2 were considered but not adopted because the required output is a task-specific tissue map trained from labeled wound-tissue data.",
    "Removed a duplicated malformed sentence and stated the model actually used.",
)
replace(
    "Chapter 2 – Wound Segmentation",
    "Chen et al. (2023) investigated a diabetic-foot-ulcer segmentation strategy that combines HarDNet-MSEG and Segment Anything Model components. The study demonstrates the value of combining a task-specific segmentation architecture with a general segmentation foundation model. For Wound-Splat, the segmentation output creates a clinically interpretable wound mask and tissue map, and restricts the reconstructed point cloud to the wound region before measurement.",
    "Chen et al. (2023) investigated a diabetic-foot-ulcer segmentation strategy that combines HarDNet-MSEG and Segment Anything Model components. The study demonstrates the value of task-specific wound segmentation. In Wound-Splat, three-dimensional wound isolation and two-dimensional tissue characterization are separate mechanisms: the reconstructed point cloud is isolated through geometric and color-based filtering, while the trained segmentation model produces a tissue map for an operator-selected image region.",
    "Corrected the statement that the 2D tissue model restricts the 3D point cloud.",
)
replace(
    "Chapter 2 – Wound Segmentation",
    "Two limitations follow from this design. The dataset defines only three tissue classes, so necrotic tissue, eschar, and slough have no corresponding label and would be assigned one of the three available categories. The model also has no \"not a wound\" class, so the region submitted for analysis must be selected to contain a wound; the system reports when no wound tissue is found rather than forcing a label.",
    "Two limitations follow from this design. The dataset provides background plus three labeled tissue classes: fibrin, granulation, and callus. Tissue appearances outside these labels, including unsupported or ambiguous categories, may be assigned to one of the available classes. Because the training images are tightly cropped wound images, the background class is not an out-of-distribution or general \"not a wound\" detector. The operator must therefore select a suitable wound region. When predicted non-background wound pixels occupy less than one percent of the selected region, the system returns a no-wound result and withholds tissue percentages.",
    "Corrected the class description and documented the actual one-percent no-wound rule.",
)
replace(
    "Chapter 2 – ISO/IEC 25010",
    "A wound-assessment application is not evaluated only by its algorithmic accuracy. Users must be able to upload the intended video, understand the result, and repeat the procedure consistently. The ISO/IEC 25010:2023 product-quality model provides a structured vocabulary for evaluating software quality, including functional suitability, performance efficiency, compatibility, usability, reliability, security, maintainability, and portability (International Organization for Standardization, 2023). The Wound-Splat study selects functional suitability, performance efficiency, reliability, security, and usability as evaluation dimensions, which should be defined with observable indicators and user-rating criteria.",
    "A wound-assessment application is not evaluated only by its algorithmic accuracy. Users must be able to submit the intended video, understand the result, and repeat the procedure consistently. ISO/IEC 25010:2023 provides a product-quality model with nine characteristics: functional suitability, performance efficiency, compatibility, interaction capability, reliability, security, maintainability, flexibility, and safety (International Organization for Standardization, 2023). For the planned respondent evaluation, Wound-Splat selects functional suitability, performance efficiency, interaction capability, reliability, and safety. Objective geometric and segmentation accuracy are evaluated separately through technical tests rather than respondent opinion.",
    "Updated the ISO/IEC 25010:2023 terminology and separated objective accuracy tests from user ratings.",
)
replace(
    "Chapter 2 – Synthesis and Research Gap",
    "The literature also shows that segmentation and reconstruction should not be developed as independent modules. A wound mask can remove irrelevant background points, improve the allocation of Gaussian primitives, and provide the boundary required for measurement. Conversely, a multi-view or three-dimensional reconstruction can improve the consistency of segmentation across frames. Datasets such as DFUC2022, Syn3DWound, and SALVE provide useful benchmarks, but phantom-based validation and public-image segmentation do not fully reproduce the conditions of longitudinal clinical monitoring.",
    "The literature shows that segmentation and reconstruction address related but distinct requirements. In the present system, the three-dimensional wound region is isolated from the reconstructed point cloud for geometric measurement, while a separate two-dimensional model characterizes tissue within an operator-selected image region. Future multi-view or three-dimensional segmentation could improve consistency across frames. Datasets such as DFUC2022, Syn3DWound, and SALVE provide useful benchmarks, but synthetic, phantom, and public-image tests do not fully reproduce longitudinal clinical monitoring.",
    "Aligned the synthesis with the actual separation between 3D isolation and 2D tissue analysis.",
)
replace(
    "Chapter 2 – Synthesis and Research Gap",
    "The central research gap is therefore not the absence of smartphone wound imaging or the absence of 3D reconstruction in isolation. More specifically, there remains a need for an integrated and explicitly evaluated workflow that combines smartphone video input, SfM-based camera estimation, wound-region segmentation, fast 3DGS reconstruction, geometric surface generation, volumetric measurement, registration across visits, and clinically interpretable reporting. Wound3DAssist demonstrates that integrated 3D wound assessment is feasible, so the contribution of Wound-Splat should be framed as an investigation of a 3DGS-centered, GPU-backed, longitudinal DFU monitoring workflow and its performance under the study’s defined phantom and software-quality evaluation conditions.",
    "The central research gap is therefore not the absence of smartphone wound imaging or three-dimensional reconstruction in isolation. There remains a need for an explicitly evaluated workflow connecting smartphone video capture, SfM camera estimation, interactive 3DGS visualization, wound-region isolation, physical-scale estimation, traceable geometric measurement, tissue characterization, and understandable reporting. Wound3DAssist demonstrates that integrated smartphone-based 3D wound assessment is feasible, but performance cannot be assumed to transfer to a different reconstruction pipeline, capture protocol, measurement method, or hardware configuration.",
    "Focused the research gap on the implemented workflow and removed an implication that longitudinal monitoring is already operational.",
)
replace(
    "Chapter 2 – Synthesis and Research Gap",
    "Wound-Splat addresses this gap by integrating these stages in a client-server application and by mapping measurable outputs to the Wagner and University of Texas systems while acknowledging their clinical limitations. Its contribution lies in the system-level integration, the use of smartphone video and 3DGS for accessible reconstruction, the use of segmentation to isolate the wound region, the conversion of reconstructed points into a measurable surface, and the comparison of wound models across visits. The framework should not be presented as a substitute for clinical diagnosis; rather, it is a technical decision-support system whose accuracy, reliability, usability, security, and limitations must be documented through transparent validation.",
    "Wound-Splat addresses this gap through the implemented integration of smartphone-video input, COLMAP Structure from Motion, stock 3D Gaussian Splatting, filtered point-cloud isolation, reference-plane and heightfield measurement, operator-guided two-dimensional tissue characterization, interactive visualization, and report generation. Synthetic point-cloud tests verify the measurement calculations, while controlled physical end-to-end validation remains necessary to establish how accurately the complete pipeline preserves wound dimensions. Longitudinal registration and Wagner and University of Texas classification support remain Thesis 2 objectives. The framework is a technical decision-support prototype and not a substitute for clinical diagnosis.",
    "Corrected the claimed current contributions and stated the remaining validation boundary.",
)

replace(
    "Chapter 2 – Related Systems",
    "Wound3DAssist provides a particularly close comparison to the proposed system. Chierchia et al. (2025b) presented a framework using short monocular consumer-grade videos and integrated 3D reconstruction, wound segmentation, tissue classification, and periwound analysis. The record reports evaluation on digital models, silicone phantoms, and real patients, with millimeter-level accuracy and assessment times below twenty minutes. Wound-Splat differs in its intended emphasis on a 3DGS-centered pipeline, GPU-backend deployment, longitudinal registration, and mapping to the Wagner and University of Texas systems.",
    "Wound3DAssist provides a particularly close comparison to the proposed system. Chierchia et al. (2025b) presented a framework using short monocular consumer-grade videos and integrated 3D reconstruction, wound segmentation, tissue classification, and periwound analysis. The record reports evaluation on digital models, silicone phantoms, and real patients, with millimeter-level accuracy and assessment times below twenty minutes. Wound-Splat differs in its implemented emphasis on a 3DGS-centered, GPU-backed reconstruction and measurement pipeline. Longitudinal registration and Wagner and University of Texas classification support remain planned Thesis 2 extensions.",
    "Kept the comparison but removed wording that presented planned Thesis 2 functions as implemented features.",
)
replace(
    "Chapter 2 – Surface Reconstruction and Geometric Wound Measurement",
    "A wound-volume calculation requires a surface or volume representation with a meaningful boundary. A point cloud or collection of Gaussian primitives may be sufficient for visualization but is not automatically a closed surface. Wound-Splat therefore separates these two requirements. A surface-reconstruction stage converts the segmented points into a triangle mesh using Poisson surface reconstruction, which provides the smooth, closed surface used for visualization and for exporting the wound model. Geometric measurement, however, is performed directly on the segmented point cloud rather than on the mesh, so that the reported values do not inherit the smoothing and hole-filling behavior that Poisson reconstruction applies to concave regions.",
    "A wound-volume calculation requires a surface or volume representation with a meaningful boundary. A point cloud or collection of Gaussian primitives may be sufficient for visualization but is not automatically a closed surface. Wound-Splat therefore separates visualization from measurement. An optional surface-reconstruction stage can convert the isolated points into a triangle mesh using Poisson surface reconstruction for visualization and export. Geometric measurement is performed directly on the isolated point cloud rather than on the mesh so that reported values do not inherit the smoothing and hole-filling behavior that Poisson reconstruction may apply to concave regions.",
    "Clarified that Poisson reconstruction is optional and does not guarantee a closed measurement surface.",
)
replace(
    "Chapter 2 – Wound Isolation",
    "Wound isolation. Wound-Splat isolates the wound after reconstruction rather than before it. Wound tissue is strongly saturated and red, while dressings, reference markers, and background are not, so points in the reconstructed cloud are gated by hue and saturation and the largest connected cluster is retained. A margin of surrounding skin is deliberately preserved, because the measurement stage fits its reference plane to that periwound region. The method is fast and requires no training data, but it depends on color and is therefore sensitive to lighting, exudate, and camera color response.",
    "Wound isolation. Wound-Splat isolates the wound after reconstruction rather than before it. The point cloud is filtered by opacity and statistical outlier removal, then grouped using scale-adaptive spatial clustering. When sufficient saturated red tissue-color evidence is present, the cluster with the strongest tissue support is selected; otherwise, the largest valid cluster is used. A margin of surrounding skin is retained because the measurement stage fits its reference plane to the periwound region. The method is fast and requires no labeled three-dimensional data, but it remains sensitive to lighting, exudate, camera color response, and reconstruction noise.",
    "Matched the description to the implemented point-cloud filtering and cluster-selection logic.",
)

# Definition of Terms table: add a header and correct four misleading entries.
terms = doc.tables[0]
hdr = OxmlElement("w:tr")
for text in ("Term", "Operational Definition"):
    tc = OxmlElement("w:tc")
    tcPr = OxmlElement("w:tcPr")
    tc.append(tcPr)
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    b = OxmlElement("w:b")
    rPr.append(b)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    p.append(r)
    tc.append(p)
    hdr.append(tc)
terms._tbl.insert(0, hdr)
log("Chapter 2 – Definition of Terms", "No header row", "Added 'Term' and 'Operational Definition' header row", "Improved table readability without changing the table structure.")

term_updates = {
    "Longitudinal Monitoring": "Planned Thesis 2 comparison of the same wound across visits after geometric registration and quality checking.",
    "Poisson Surface Reconstruction": "An optional visualization method that converts oriented points into a smooth mesh; its output is not used as the validated measurement surface.",
    "Wagner and University of Texas Systems": "Clinical classification systems that require findings beyond geometry alone, including infection, ischemia, and anatomical involvement; Wound-Splat classification support remains planned for Thesis 2.",
    "Scale Calibration": "Conversion from reconstruction units to physical units using a supported bank card or coin visible in the captured scene.",
    "Temporal Tracking": "A planned Thesis 2 module for aligning wound models across visits and calculating changes in measurements.",
}
for row in terms.rows[1:]:
    key = row.cells[0].text.strip()
    if key in term_updates:
        old = row.cells[1].text.strip()
        set_cell_text(row.cells[1], term_updates[key])
        log(f"Chapter 2 – Definition of Terms: {key}", old, term_updates[key], "Aligned the definition with the current implementation boundary.")

# Chapter 3: dependency and data corrections.
for table in doc.tables:
    if table.rows and table.rows[0].cells[0].text.strip() == "Package":
        for row in table.rows[1:]:
            package = row.cells[0].text.strip()
            if package == "segmentation-models-pytoch":
                set_cell_text(row.cells[0], "segmentation-models-pytorch")
                log("Chapter 3 – Table 4", package, "segmentation-models-pytorch", "Corrected the package name.")
            if package == "pytest" and row.cells[1].text.strip() == "0.1.1":
                set_cell_text(row.cells[1], "9.1.1")
                log("Chapter 3 – Table 4", "pytest 0.1.1", "pytest 9.1.1", "Corrected the installed test-runner version.")
            if package == "pypdf" and "Verification od" in row.cells[2].text:
                set_cell_text(row.cells[2], "Verification of generated tissue-report content")
                log("Chapter 3 – Table 4", "Verification od generated tissue report content", "Verification of generated tissue-report content", "Corrected a typographical error.")

replace(
    "Chapter 3 – Hardware",
    "\tA rented cloud processor handles runs that exceed that limit. 3DGS optimization holds every Gaussian primitive in the video memory, and this requirement grows with the number of frames and the density of the scene. The rented instance is providing 48GB, runs the full stack under the headless Linux operating system, and is billed hourly, so it is provisioned only when needed.",
    "A compatible cloud GPU instance may be provisioned manually for selected high-memory reconstruction runs. The current application does not automatically detect an out-of-memory condition, transfer a queued job to the cloud, or provide automatic failover. When the cloud environment is used, the full processing stack is started separately on the headless Linux instance.",
    "Corrected the cloud workflow to reflect manual provisioning rather than automatic job handling.",
)
replace(
    "Chapter 3 – Dataset",
    "\tPermission and terms of use are available. DFUTisse is not openly licensed. Written permission was granted by Prof. Zeyun Yu of the University of Wisconsin-Milwaukee on September 4, 2026, restricting use to nonprofit academic research, prohibiting redistribution, and requiring citation. The dataset is therefore excluded from the project's public repository and not distributed anywhere within the software; the permission correspondence is retained as a record of compliance.",
    "Permission and terms of use are documented. DFUTissue is not openly licensed. Written permission was granted by Prof. Zeyun Yu of the University of Wisconsin-Milwaukee on September 4, 2026, restricting use to nonprofit academic research, prohibiting redistribution, and requiring citation. The dataset is excluded from the public repository and is not distributed with the software; the permission correspondence is retained as a record of compliance.",
    "Corrected the dataset name and tightened the permission statement.",
)

data_sources = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Source" and "Quantity" in t.rows[0].cells[1].text)
for row in data_sources.rows[1:]:
    if row.cells[0].text.strip() == "Self-captured smartphone videos":
        old = row.cells[1].text.strip()
        set_cell_text(row.cells[1], "Multiple development recordings (non-patient)")
        log("Chapter 3 – Table 7", old, "Multiple development recordings (non-patient)", "Removed an unresolved count placeholder while preserving the limited methodological use of the recordings.")

replace(
    "Chapter 3 – Project Development",
    "One rule will govern the cycle and distinguish it from a linear sequence: an iteration is not complete when the code works, but rather when the result has been measured and a decision has been taken. The measurement may show that the addition performs adequately, in which case it is retained. It may show it performs inadequately, in which case it is revised. It may also show the addition does not justify its cost, in which case it is removed. All three outcomes occurred during development and are documented. The Wound-Splat framework development is executed through the following iterative phases:",
    "One rule governs the cycle and distinguishes it from a linear sequence: an iteration is not complete when the code runs, but when its result has been checked and a decision has been recorded. A component may be retained, revised, or removed according to the available evidence. The Wound-Splat framework was developed through the following iterative phases:",
    "Removed an unsupported statement that every possible iteration outcome had been documented.",
)
replace(
    "Chapter 3 – Project Design",
    "Client layer. All user interaction will occur in a standard web browser. The patient portal allows the patient to sign in with their provided assigned patient code, upload a wound video, observe processing status, and view the completed results. The clinical administration view will allow the staff to register patients and monitor their processing queue. The viewer will render the reconstructed model interactively in the browser, so a completed scan can be rotated in any axis and inspected without installing a specialized software. No reconstruction, measurement, or model inference occurs in this layer.",
    "Client layer. All user interaction occurs in a standard web browser. The patient portal uses an assigned patient code to look up the corresponding prototype record, upload a wound video, observe processing status, and view completed results. This patient-code lookup is not secure authentication. The clinical administration view allows staff to register patients and monitor the processing queue. The viewer renders the reconstructed model interactively in the browser so a completed scan can be rotated and inspected without specialized desktop software. No reconstruction, measurement, or model inference occurs in this layer.",
    "Corrected the patient-code workflow so it is not described as secure sign-in and improved the surrounding grammar.",
)
replace(
    "Chapter 3 – Database Design",
    "Measurements and Tissue Results. The application expects a single current measurement result and a single current tissue result for each scan. The Measurement table stores surface area, volume, maximum depth, width, and height. However, its scan_id foreign key is not currently unique at the database level, meaning the schema can technically contain more than one measurement row for a scan. The TissueResult table enforces a unique scan_id and therefore permits only one current tissue result for each scan. Repeating tissue analysis updates the existing result rather than creating a history of previous analyses.",
    "Measurements and Tissue Results. The application expects a single current measurement result and a single current tissue result for each scan. The Measurement table stores planar wound opening area in the surface-area field, together with cavity volume, maximum depth, width, and height. However, its scan_id foreign key is not currently unique at the database level, meaning the schema can technically contain more than one measurement row for a scan. The TissueResult table enforces a unique scan_id and therefore permits only one current tissue result for each scan. Repeating tissue analysis updates the existing result rather than creating a history of previous analyses.",
    "Clarified that the database surface-area field contains the implemented planar wound opening area.",
)

# Fix the combined database heading/image/caption without losing the embedded image.
dbp = next(p for p in doc.paragraphs if p.text.strip() == "Database DesignFigure 9: Database Entity Relationship Diagram")
dbp.insert_paragraph_before("Database Design", style="Heading 3")
for r in dbp.runs:
    if r.text:
        r.text = ""
dbp.style = doc.styles["normal"]
caption = OxmlElement("w:p")
dbp._p.addnext(caption)
from docx.text.paragraph import Paragraph
caption_para = Paragraph(caption, dbp._parent)
caption_para.style = doc.styles["Heading 5"]
caption_para.add_run("Figure 9: Database Entity Relationship Diagram")
log("Chapter 3 – Database Design", "Database Design and Figure 9 caption merged in one paragraph", "Separated the heading, diagram, and Figure 9 caption", "Corrected the layout while preserving the embedded diagram.")

# Insert the missing Algorithm Implementation section before the constraints section.
anchor = find_paragraph("Multiple Constraints and Trade-off Analysis")
algorithm_parts = [
    ("Algorithm Implementation", "Heading 2", "Chapter 3 – Algorithm Implementation", "Added the missing implementation section requested by the authors."),
    ("The implemented algorithm connects the web application to a seven-stage reconstruction and measurement pipeline. The FastAPI backend stores the uploaded video and scan record, places the scan identifier in an in-process queue, and uses one background worker to process queued reconstruction jobs sequentially. The frontend receives the scan identifier immediately and polls the stored progress fields while processing continues.", "normal", "Chapter 3 – Algorithm Implementation", "Documented the actual queue and progress behavior."),
    ("Step 1 – Frame Extraction", "Heading 3", "Chapter 3 – Algorithm Implementation", "Documented the first implemented pipeline stage."),
    ("FFmpeg samples the uploaded video at two frames per second and writes sequential JPEG images to the scan input directory. The fixed sampling rate limits the number of images passed to reconstruction while preserving overlapping viewpoints when the capture protocol is followed.", "normal", "Chapter 3 – Algorithm Implementation", "Recorded the implemented two-frames-per-second setting."),
    ("Step 2 – Camera-Pose Estimation and Sparse Reconstruction", "Heading 3", "Chapter 3 – Algorithm Implementation", "Documented the second implemented pipeline stage."),
    ("COLMAP processes the extracted frames to detect and match visual features, estimate camera poses, and generate the sparse scene model required by 3DGS. The application stores the numbers of extracted and registered frames and computes the registration rate. Failure to produce the required camera-pose output stops the scan and records the error.", "normal", "Chapter 3 – Algorithm Implementation", "Documented the COLMAP inputs, outputs, and failure behavior."),
    ("Step 3 – 3D Gaussian Splatting Optimization and Preview Rendering", "Heading 3", "Chapter 3 – Algorithm Implementation", "Documented the optimization settings and visualization output."),
    ("The stock 3D Gaussian Splatting implementation uses the COLMAP output as initialization. The Gaussian model is optimized for 30,000 iterations with a resolution factor of 2, then the final saved representation is rendered to produce preview images. The Gaussian representation supports interactive viewing, but it is not treated as the validated measurement surface.", "normal", "Chapter 3 – Algorithm Implementation", "Recorded the actual iteration and resolution settings and the role of the splat representation."),
    ("Step 4 – Three-Dimensional Wound Isolation", "Heading 3", "Chapter 3 – Algorithm Implementation", "Documented the wound-isolation algorithm."),
    ("The final point cloud is filtered by Gaussian opacity and statistical outlier removal. A scale-adaptive DBSCAN procedure groups the remaining points. When sufficient saturated red tissue evidence is available, the cluster with the strongest tissue-color support is selected; otherwise, the largest valid cluster is used. A surrounding periwound margin is retained because the measurement algorithm requires nearby skin points for reference-plane estimation. The isolated cloud is saved separately from the optional viewer-oriented splat filtering.", "normal", "Chapter 3 – Algorithm Implementation", "Described the implemented filtering, clustering, color evidence, and periwound margin."),
    ("Step 5 – Scale Calibration and Geometric Measurement", "Heading 3", "Chapter 3 – Algorithm Implementation", "Documented the calibration and measurement algorithm."),
    ("When the user selected a supported bank card or coin as the reference type, the calibration routine estimates centimeters per reconstruction unit from the captured reference. Failure to obtain a reliable scale does not automatically fail the scan; the resulting size-dependent values remain uncalibrated. For measurement, statistical outliers are removed, a periwound reference plane is fitted with RANSAC, and point depths are computed relative to that plane. Cavity points are projected to the plane and rasterized into a median depth heightfield. Internal gaps are filled, the largest connected footprint is retained, and the system calculates planar wound opening area, cavity volume, maximum depth, and PCA-derived width and height. The optional Poisson mesh is used for visualization only.", "normal", "Chapter 3 – Algorithm Implementation", "Documented the implemented RANSAC, heightfield, calibration, and metric calculations."),
    ("Step 6 – Tissue Characterization", "Heading 3", "Chapter 3 – Algorithm Implementation", "Documented the separate tissue-analysis process."),
    ("After reconstruction, the operator opens a selected frame and draws a rectangular region containing the wound. The region is resized to 256 × 256 pixels and processed by a U-Net-style model with a MiT-b3 encoder, scSE attention, and four output classes: background, fibrin, granulation, and callus. The predicted class mask is resized back to the selected region. Tissue percentages use predicted non-background pixels as the denominator. If fewer than one percent of the selected pixels are predicted as wound tissue, the system records a no-wound result and withholds the three percentages. The selected frame, region, overlay, wound fraction, and result are stored for traceability.", "normal", "Chapter 3 – Algorithm Implementation", "Documented the actual model, preprocessing, denominator, threshold, and stored outputs."),
    ("Step 7 – Persistence, Visualization, and Reporting", "Heading 3", "Chapter 3 – Algorithm Implementation", "Documented final output handling."),
    ("Structured patient, scan, measurement, and tissue-result data are stored in SQLite, while uploaded videos and generated files remain in scan-specific file-system directories. The browser retrieves the available point cloud, optional mesh, or Gaussian representation for interactive viewing. ReportLab generates the initial PDF after geometric measurement, and the report is rebuilt after successful tissue analysis. Required-stage exceptions mark the scan as failed and allow the worker to continue to the next queued scan. Longitudinal registration and Wagner and University of Texas classification support are not part of this implemented algorithm and remain planned for Thesis 2.", "normal", "Chapter 3 – Algorithm Implementation", "Documented storage, viewing, reporting, error handling, and the Thesis 2 boundary."),
]
for text, style, location, reason in algorithm_parts:
    insert_before(anchor, text, style, location, reason)

# Formula and trade-off corrections.
replace(
    "Chapter 3 – Constraint Definition",
    "This section defines the criteria that will be used to evaluate the Wound-Splat framework under its hardware, reconstruction, measurement, and dataset constraints. The evaluation focuses on processing time, image-registration success, and geometric measurement error, scale calibration, tissue-segmentation performance, and pipeline reliability. These criteria will be used to determine whether the framework can produce consisten results from standard smartphone videos while operating on the available computing hardware.",
    "This section defines the criteria that will be used to evaluate the Wound-Splat framework under its hardware, reconstruction, measurement, and dataset constraints. The evaluation focuses on processing time, image-registration success, geometric measurement error, scale calibration, tissue-segmentation performance, and pipeline reliability. These criteria will be used to determine whether the framework can produce consistent results from standard smartphone videos while operating on the available computing hardware.",
    "Corrected grammar and a typographical error.",
)
replace(
    "Chapter 3 – Formula 1",
    "Queue waiting time must be recorded separately because the system processes reconstruction jobs sequentially. This prevents scans waiting in the queue from being interpreted as having longer computational processing times.",
    "Queue waiting time must be recorded separately because the system processes reconstruction jobs sequentially. The current schema stores submission and completion times but does not yet store a dedicated active-processing start time; a started-at timestamp or an equivalent controlled timing log must therefore be added before Formula 1 is used in the Thesis 2 experiment.",
    "Identified the instrumentation required to calculate active processing time correctly.",
)
replace(
    "Chapter 3 – Formula 4",
    "\tThe current synthetic-phantom test uses a maximum error tolerance of 10% for each evaluated geometric quantity. This criterion verifies the internal measurement calculations for the analytically defined point cloud. It does not validated the complete smartphone video-to-measurement pipeline.",
    "The current synthetic-phantom test uses a maximum error tolerance of 10% for each evaluated geometric quantity. MAPE will be reported separately for planar opening area, cavity volume, maximum depth, width, and height before any overall summary is calculated. This criterion verifies the measurement calculations for analytically defined point clouds and does not validate the complete smartphone video-to-measurement pipeline.",
    "Corrected grammar and specified that MAPE is calculated per metric.",
)
replace("Chapter 3 – Formula 5", "s = Lreal / Lmodel", "s = Lref,real / Lref,model", "Distinguished the reference-object dimension from a measured wound length.")
replace("Chapter 3 – Formula 5", "Lreal = known physical dimension of the reference object; and", "Lref,real = known physical dimension of the reference object; and", "Matched the variable definition to the revised equation.")
replace("Chapter 3 – Formula 5", "Lmodel = corresponding reference-object dimension measured in reconstruction units.", "Lref,model = corresponding reference-object dimension measured in reconstruction units.", "Matched the variable definition to the revised equation.")
replace("Chapter 3 – Formula 5", "Lcm = s × Lmodel", "Lcm = s × Lmeasured", "Prevented the reference-object variable from being reused for a wound measurement.")
replace("Chapter 3 – Formula 5", "Lcm = calibrated length or depth in centimeters;", "Lmeasured = wound length or depth in reconstruction units; and\nLcm = calibrated length or depth in centimeters;", "Defined the distinct wound-measurement variable used in the conversion.")
replace(
    "Chapter 3 – Formulas 6 and 7",
    "Dice and intersection over union will be reported for the available tissue classes. Their interpretation must consider that the DFUTissue contains only 110 labeled images and does not provide labels for every tissue appearance that may occur in clinical wounds.",
    "Dice and Intersection over Union will be calculated from pooled test-set pixels for background, fibrin, granulation, and callus. The reported mean tissue score will exclude background and will omit a tissue class only when it is absent from both the prediction and the ground truth. Interpretation must consider that DFUTissue contains only 110 labeled images and does not represent every tissue appearance found in clinical wounds.",
    "Specified the actual averaging and absent-class rules used by the implementation.",
)
replace(
    "Chapter 3 – Design Trade-offs",
    "\tThe selected pipeline combines COLMAP Structure-from-Motion with 3D Gaussian Splatting. COLMAP estimates the camera positions and produces the sparse initialization required by the 3DGS optimization. The Guassian-splat representation is then used in the browser viewer because it supports interactive inspection from different viewpoints.",
    "The selected pipeline combines COLMAP Structure from Motion with 3D Gaussian Splatting. COLMAP estimates the camera positions and produces the sparse initialization required by 3DGS optimization. The Gaussian-splat representation is used in the browser viewer because it supports interactive inspection from different viewpoints.",
    "Corrected spelling and terminology.",
)
replace(
    "Chapter 3 – Design Trade-offs",
    "\tThe point-cloud heightfield was selected as the measurement representation because the optional POssing mesh may smooth the wound boundary or artificially close gaps in the reconstructed surface. The heightfield estimates the wound cavity relative to a reference plane fitted to the surrounding periwound region. However, this method assumes that the wound surface can be represented by one depth value for each grid location.",
    "The point-cloud heightfield was selected as the measurement representation because the optional Poisson mesh may smooth the wound boundary or artificially close gaps in the reconstructed surface. The heightfield estimates the wound cavity relative to a reference plane fitted to the surrounding periwound region. However, this method assumes that the wound surface can be represented by one depth value for each grid location.",
    "Corrected the Poisson spelling.",
)
replace(
    "Chapter 3 – Design Trade-offs",
    "For two-dimensional tissue characterization, the system uses an operator-selected rectangular region. This allows the clinician to identify the relevant wound region without requiring a separate automatic wound-localization model. However, the resulting tissue percentages may change when the selected frame or rectangular region changes, the selected frame, region, predicted wound fraction and the tissue percentages must therefore be retained with the result.",
    "For two-dimensional tissue characterization, the system uses an operator-selected rectangular region. This allows the operator to identify the relevant wound region without requiring a separate automatic wound-localization model. Because the resulting tissue percentages may change with the selected frame or region, the selected frame, region coordinates, predicted wound fraction, and tissue percentages must be retained with the result.",
    "Corrected a run-on sentence and avoided assuming the operator is necessarily a clinician.",
)
replace(
    "Chapter 3 – Trade-off Analysis",
    "\tThe criteria will include geometric measurement error, processing time, peak GPU-memory consumption, registration rate, completion rate, Dice, and Intersection over Union are considered benefit criteria because higher values are preferred.",
    "The cost criteria are geometric measurement error, processing time, and peak GPU-memory consumption, for which lower values are preferred. The benefit criteria are registration rate, pipeline completion rate, Dice, and Intersection over Union, for which higher values are preferred.",
    "Corrected the criterion classification and sentence structure.",
)
replace(
    "Chapter 3 – Trade-off Analysis",
    "The configuration with the highest observed value receives a rank of 10.",
    "The configuration with the highest observed value receives a rank of 10. Tied best values receive the same rank. If all compared values are zero or a configuration fails to produce a valid value, the result will be reported as not applicable or failed rather than entered into a formula with an undefined denominator.",
    "Added rules for ties, zeros, and failed configurations.",
)
replace(
    "Chapter 3 – Sensitivity Analysis",
    "\tEach criterion will be assigned a criterion weight. The sum of the weights for one trail equals 1.00. The criterion score will be obtained by dividing the criterion rank by 10 so that all criterion scores range from 0 to 1.",
    "Each criterion will be assigned a criterion weight, and the weights for one trial will sum to 1.00. Before weighting, multiple measures within one category will be averaged only after they have been normalized to the same 0-to-1 range. The criterion score will be obtained by dividing the criterion rank by 10.",
    "Corrected 'trail' and defined how multiple measures are aggregated before category weighting.",
)
replace(
    "Chapter 3 – Sensitivity Analysis",
    "Trial 5 assigns equal importance to all criteria. It provides a balanced assessment without favoring accuracy, processing time, resource consumption, or reliability.",
    "Trial 5 assigns equal importance to all criteria. It provides a balanced assessment without favoring accuracy, processing time, resource consumption, or reliability. The proposed weights in Table 12 must be reviewed with the thesis adviser before the controlled Thesis 2 evaluation and must not be changed after results are examined merely to favor a preferred configuration.",
    "Marked the trial weights as proposed and added a pre-results approval rule.",
)

# Expand the previously incomplete testing and operating procedure using the existing content.
testing_heading = find_paragraph("Testing and Operating Procedure")
old_testing1 = find_paragraph("The Wound-Splat framework is designed to operate within a client-server environment, requiring distinct operational procedures for the patient and the clinical administrator. The patient initiates the workflow by accessing the patient portal via a standard web browser, authenticating with an assigned patient code, and uploading a wound video. The upload protocol mandates that the video includes a visible scale reference (e.g., a standard bank card) and captures the wound from multiple overlapping angles. Upon upload, the system immediately returns a tracking identifier. The patient can subsequently poll the system to monitor the seven-stage processing pipeline. Once the pipeline reaches the \"rendered\" status, the patient can access the interactive 3D viewer to inspect the Gaussian splat model and download the generated PDF assessment report.")
set_paragraph_text(old_testing1, "Testing Procedure")
old_testing1.style = doc.styles["Heading 3"]
log("Chapter 3 – Testing and Operating Procedure", "Two-paragraph combined procedure", "Separated Testing Procedure and Operating Procedure with defined checks", "Completed the previously abbreviated procedure section.")
old_testing2 = find_paragraph("Concurrently, the clinical administrator operates the backend dashboard to manage patient registrations, monitor the serialized GPU processing queue, and perform tissue characterization. Because tissue analysis requires human-in-the-loop region selection, the clinician accesses the rendered scan, draws a boundary around the wound on a selected 2D frame, and triggers the tissue segmentation inference. The system then regenerates the final PDF report to include the tissue percentages alongside the geometric measurements.")
set_paragraph_text(old_testing2, "The technical evaluation will use a separate test set or controlled test inputs that were not used to tune the component being evaluated. Reconstruction tests will record extracted and registered frames, active processing time, peak GPU memory, completion status, calibration status, and required outputs. Geometric tests will compare planar opening area, cavity volume, maximum depth, width, and height with known analytical or independently measured references. Tissue-model testing will use the official labeled DFUTissue test split and report pooled per-class Dice and Intersection over Union, with the background-excluded mean tissue score. Workflow tests will verify uploads, queue progression, stored records, viewer outputs, tissue-result persistence, report generation, and required-stage error handling. Synthetic point-cloud tests will be reported separately from physical end-to-end tests because they bypass video capture, COLMAP, 3DGS, wound isolation, and automatic scale detection.")
for text, style in [
    ("Operating Procedure", "Heading 3"),
    ("1. System setup. Start the FastAPI backend, Next.js frontend, SQLite database, required reconstruction tools, and compatible NVIDIA GPU environment. A cloud GPU may be started manually when selected for an experiment; it is not an automatic fallback.", "normal"),
    ("2. Patient and scan setup. Staff create or select the appropriate patient record and assign a unique patient code. The code supports prototype record lookup and must not be treated as secure authentication.", "normal"),
    ("3. Capture and submission. Record the subject under controlled lighting from overlapping viewpoints while keeping the supported card or coin visible. Select the reference type, upload the video, and retain the returned scan identifier.", "normal"),
    ("4. Processing and review. Monitor the seven-stage status display while the in-process worker handles queued scans sequentially. When processing finishes, review the registration rate, calibration status, previews, interactive representation, measurements, and PDF report. A missing scale produces an uncalibrated result rather than automatically failing the scan.", "normal"),
    ("5. Tissue analysis. Open the tissue-analysis interface, use the selected frame, draw a tight rectangular region around the wound, run inference, and review the overlay, wound fraction, and available tissue percentages. The system rebuilds the report after a successful tissue result.", "normal"),
    ("6. Error and feedback handling. Record failed required stages, preserve the error message, and confirm that the queue continues to the next scan. Document operator observations about capture quality, clarity, and workflow behavior for later refinement. Automated measurements remain decision-support information and require clinical interpretation.", "normal"),
]:
    new_p = OxmlElement("w:p")
    old_testing2._p.addnext(new_p)
    para = Paragraph(new_p, old_testing2._parent)
    para.style = doc.styles[style]
    para.add_run(text)
    old_testing2 = para

# Project Evaluation: current ISO terms, objective tests, and no unimplemented claims.
replace(
    "Chapter 3 – Project Evaluation",
    "The evaluation design of the system is grounded in the ISO/IEC 25010:2023 Systems and Software Quality Requirements and Evaluation (SQuaRE) model, which provides a standardized framework for assessing software product quality. This internationally recognized standard evaluates the functionality, reliability, usability, and overall quality of the system, making it ideal for a healthcare-focused framework that processes sensitive wound imagery and generates clinical decision-support metrics.",
    "The planned software-quality evaluation is guided by ISO/IEC 25010:2023. The respondent instrument uses functional suitability, performance efficiency, interaction capability, reliability, and safety because these characteristics can be assessed through observable use of the prototype. Objective geometric accuracy, tissue-segmentation performance, processing time, GPU-memory use, registration rate, and pipeline completion rate will be measured through the technical procedures defined earlier and will not be determined from respondent opinion.",
    "Updated the standard terminology and separated respondent perception from objective technical evaluation.",
)
replace(
    "Chapter 3 – Project Evaluation",
    "Functional Suitability. The system is designed to address the specific needs of both diabetic patients and healthcare providers by delivering accurate, volumetric wound assessments using standard smartphone videos. It integrates 3D Gaussian Splatting (3DGS) reconstruction with 2D tissue segmentation to ensure that geometric measurements and tissue characterizations are both clinically relevant. The system also provides essential details such as wound volume, depth, surface area, and mapped Wagner and University of Texas classification suggestions. For clinicians, it serves as an accessible decision-support tool to monitor healing trajectories without requiring specialized, expensive 3D medical scanners.",
    "Functional Suitability. Respondents will assess whether the prototype provides the functions demonstrated in the defined workflow: scan submission, progress monitoring, reconstruction viewing, calibration-status display, geometric measurement display, operator-guided tissue characterization, and report access. Longitudinal comparisons and Wagner and University of Texas classification support will be excluded until implemented in Thesis 2.",
    "Removed claims about accuracy, healing trajectories, and unfinished classification features.",
)
replace(
    "Chapter 3 – Project Evaluation",
    "Functional Correctness. The framework ensures the provision of correct results and precise measurements through its rigorously tested computational pipeline. The geometric accuracy of the 3D reconstruction and the absolute scale calibration routine are validated against physical rigid objects of known dimensions and synthetic wound phantoms (spherical caps, half-ellipsoids, and shallow craters) with analytically known ground truths. By measuring the raw segmented point cloud rather than the smoothed Poisson mesh, the system minimizes geometric degradation, ensuring that the extracted volumetric and spatial metrics precisely reflect the actual wound geometry without artificial smoothing.",
    "Interaction Capability. Respondents will assess whether capture guidance, form labels, status indicators, error messages, the interactive viewer, tissue-analysis controls, measurement labels, calibration warnings, and the generated report are understandable and usable for the tasks assigned to them.",
    "Replaced a respondent-rated correctness claim with the selected ISO/IEC 25010:2023 interaction-capability characteristic.",
)
replace(
    "Chapter 3 – Project Evaluation",
    "Performance Efficiency. The system is optimized to utilize available hardware resources effectively while managing user expectations during computationally intensive tasks. Although a complete 3D Gaussian Splatting reconstruction requires significant GPU processing, the framework employs a serialized background worker queue that prevents memory exhaustion on consumer-grade hardware (e.g., 6GB VRAM). For scans that exceed local memory limits, the system dynamically utilizes the provisioned 48GB cloud instance. The application maintains performance efficiency by providing real-time progress polling to the client, ensuring that the user interface remains responsive and does not time out while the backend processes the seven-stage pipeline across the hybrid hardware environment.",
    "Performance Efficiency. Respondents will assess whether uploads, status updates, navigation, viewing, and report access remain responsive during the demonstrated workflow. Active processing time and peak GPU-memory consumption will be measured objectively on the designated hardware. The sequential worker reduces simultaneous GPU contention but does not guarantee that every scan will fit in 6 GB of video memory, and cloud execution is manually provisioned rather than automatic.",
    "Removed the unsupported automatic cloud-offload and guaranteed memory-prevention claims.",
)
replace(
    "Chapter 3 – Project Evaluation",
    "Reliability. The framework is designed to maintain a specified level of performance and availability under stated conditions, particularly when handling variable-quality smartphone inputs. If a stage in the processing pipeline fails, such as insufficient visual overlap for Structure from Motion (SfM) or failure to detect the scale reference object, the system isolates the error, marks the specific scan record as \"failed\" with a descriptive log, and safely advances the queue to the next scan. This fault-tolerant architecture ensures that a single corrupted upload or processing anomaly does not crash the background worker or disrupt the availability of the web application for other users.",
    "Reliability. Respondents will assess whether the demonstrated workflow preserves submitted scans, reports status consistently, retains completed outputs, presents understandable errors, and remains usable after a required-stage failure. Required-stage exceptions mark the scan as failed and allow the worker to continue. Failure to detect a usable reference object is treated differently: the scan may complete with uncalibrated measurements and a visible warning.",
    "Corrected the scale-reference behavior and avoided promising complete fault tolerance before testing.",
)
replace(
    "Chapter 3 – Project Evaluation",
    "Safety. The system incorporates safeguards to prevent unacceptable risks of harm to patients or data breaches, recognizing its role as a clinical decision-support tool rather than an autonomous diagnostic device. To ensure clinical safety, the framework explicitly flags uncalibrated measurements as \"approximate\" when the reference object is not detected, preventing clinicians from making critical decisions based on arbitrary model scaling. Furthermore, the system is designed to handle sensitive health information securely, ensuring that patient data and wound imagery are stored in a protected relational database with strict access controls, and all generated outputs clearly distinguish automated algorithmic suggestions from confirmed clinical diagnoses.",
    "Safety. Respondents will assess whether the interface and report identify the prototype as decision support, display calibration status, distinguish automated outputs from confirmed clinical findings, and avoid presenting planned longitudinal or classification functions as completed results. The present patient-code lookup is not secure authentication or role-based authorization; privacy and security controls require separate technical and institutional evaluation before use with patient data.",
    "Removed unsupported security claims and stated the current prototype limitation.",
)

# Questionnaire correction: five-point scale, correct respondent heading, and observable statements.
first_survey_heading = [p for p in doc.paragraphs if p.text.strip() == "Survey Questionnaire for Non-Information Technology Respondents"][0]
set_paragraph_text(first_survey_heading, "Survey Questionnaire for Information Technology Respondents")
log("Chapter 3 – Survey Questionnaires", "First questionnaire incorrectly labeled Non-Information Technology", "Survey Questionnaire for Information Technology Respondents", "Corrected the duplicated respondent-group title.")
second_heading = [p for p in doc.paragraphs if p.text.strip() == "Survey Questionnaire for Non-Information Technology Respondents"][0]
scale_para = first_survey_heading.insert_paragraph_before("Response scale for both questionnaires: 5 – Strongly Agree; 4 – Agree; 3 – Neutral; 2 – Disagree; 1 – Strongly Disagree.")
scale_para.style = doc.styles["normal"]
log("Chapter 3 – Survey Questionnaires", "Six-point rating columns without verbal anchors", "Five-point Likert scale with Strongly Agree to Strongly Disagree anchors", "Matched the stated project-evaluation format and made the scale interpretable.")

survey_tables = doc.tables[12:22]
if len(survey_tables) != 10:
    raise RuntimeError(f"Expected 10 survey tables, found {len(survey_tables)}")

it_content = [
    ("Functional Suitability - Assesses whether the implemented functions support the defined scan, measurement, tissue-analysis, and reporting workflow.", [
        "1. The system accepts a valid scan submission and produces the expected reconstruction, measurement, and report outputs when required stages complete.",
        "2. Calibration status and geometric measurements are stored and displayed consistently across the scan record, viewer, and report.",
        "3. Tissue analysis stores the selected region and produces an overlay, wound fraction, and tissue percentages or a no-wound result as defined.",
    ]),
    ("Interaction Capability - Assesses whether technical users can understand and operate the interface and inspect its outputs.", [
        "1. Capture guidance, field labels, and required inputs are clear.",
        "2. Progress stages, completion states, calibration warnings, and error messages are understandable.",
        "3. The viewer, tissue-analysis controls, measurements, and PDF report are organized clearly for review.",
    ]),
    ("Performance Efficiency - Assesses observable responsiveness while objective runtime and memory are recorded separately.", [
        "1. Upload submission and progress-status updates remain responsive while reconstruction runs in the background.",
        "2. Multiple submitted scans remain visible and are processed sequentially without losing queued requests.",
        "3. The viewer and report can be opened without causing browser instability on the designated client device.",
    ]),
    ("Reliability - Assesses whether the workflow preserves records, reports failures, and recovers for later scans.", [
        "1. Completed scans retain their linked measurements, files, and reports for later retrieval.",
        "2. Required-stage failures produce a clear scan error and do not prevent the worker from processing the next queued scan.",
        "3. A missing scale reference is reported as an uncalibrated result rather than being misreported as a calibrated measurement.",
    ]),
    ("Safety - Assesses whether the prototype communicates limitations that reduce the risk of misinterpretation.", [
        "1. The interface and report present automated outputs as decision-support information rather than confirmed clinical diagnoses.",
        "2. Uncalibrated measurements are visibly identified and are not presented as validated physical measurements.",
        "3. Planned longitudinal and Wagner or University of Texas outputs are not displayed as completed functions in the current prototype.",
    ]),
]

nonit_content = [
    ("Functional Suitability - Assesses whether the visible prototype functions support the assigned wound-scan workflow.", [
        "1. I can submit the required video and reference information using the web interface.",
        "2. I can view the completed reconstruction, measurement labels, calibration status, and report.",
        "3. I can complete the tissue-analysis task and understand whether percentages or a no-wound result were produced.",
    ]),
    ("Interaction Capability - Assesses the clarity and ease of use of the interface.", [
        "1. The capture and upload instructions are easy to understand.",
        "2. The progress display and error messages clearly explain the scan status.",
        "3. The viewer, tissue-analysis controls, and PDF report are easy to navigate and interpret.",
    ]),
    ("Performance Efficiency - Assesses the responsiveness experienced by the user.", [
        "1. The interface responds promptly when I submit a scan or move between pages.",
        "2. Progress updates appear often enough for me to understand that processing is continuing.",
        "3. The completed viewer and report open and remain usable on the provided device.",
    ]),
    ("Reliability - Assesses whether users can retrieve results and understand unsuccessful operations.", [
        "1. Completed scan records and reports remain available when I return to them.",
        "2. When a required processing stage fails, the system provides a clear message instead of showing an incomplete result as complete.",
        "3. A scan without reliable scale calibration is clearly identified as uncalibrated.",
    ]),
    ("Safety - Assesses whether the prototype communicates its limits clearly.", [
        "1. The system clearly states that its outputs support, but do not replace, clinical assessment.",
        "2. Calibration warnings and measurement units are visible and understandable.",
        "3. The report distinguishes automated outputs from clinician-confirmed findings.",
    ]),
]

for group_name, tables, content in [("IT", survey_tables[:5], it_content), ("Non-IT", survey_tables[5:], nonit_content)]:
    for table, (description, items) in zip(tables, content):
        if len(table.columns) == 7:
            remove_last_table_column(table)
        # Header row and scale row.
        set_cell_text(table.rows[0].cells[0], "Criteria")
        for i in range(1, 6):
            set_cell_text(table.rows[0].cells[i], "Rating")
        set_cell_text(table.rows[1].cells[0], description)
        for i, score in enumerate([5, 4, 3, 2, 1], start=1):
            set_cell_text(table.rows[1].cells[i], str(score))
        for row, item in zip(table.rows[2:5], items):
            set_cell_text(row.cells[0], item)
    log(f"Chapter 3 – {group_name} questionnaire", "Items asked respondents to judge algorithmic accuracy, security controls, GPU use, or unimplemented functions", "Revised five ISO-aligned tables with observable, role-appropriate statements", "Prevented respondents from being used as substitutes for objective technical or security evaluation.")

# Commercialization: conditional future language and actual execution model.
replace(
    "Chapter 3 – Potential for Commercialization",
    "The system aims to revolutionize chronic wound management by transitioning from subjective 2D photographic assessments to objective 3D volumetric tracking. Utilizing 3D Gaussian Splatting (3DGS), the framework automates the extraction of precise wound depth, surface area, and cavity volume; metrics that are critical for predicting DFU healing trajectories but impossible to capture via standard smartphone photos (Kerbl et al., 2023). Furthermore, the integration of an AI-driven tissue segmentation model automates the classification of wound etiologies (fibrin, granulation, and callus), minimizing the need for manual, subjective diagnosis by overworked clinical staff (Dhar et al., 2024). Because the heavy computational lifting is offloaded to a cloud-based serialized queue, the web-based application provides patients in resource-limited settings access to advanced 3D monitoring using only standard consumer smartphones, drastically reducing the financial and logistical burden of frequent in-person clinic visits.",
    "If controlled physical and clinical evaluation demonstrates acceptable accuracy, reliability, safety, and usability, Wound-Splat may support more objective documentation of wound geometry than two-dimensional photographs alone. The current prototype estimates planar wound opening area, cavity volume, maximum depth, width, and height and provides separate tissue-characterization outputs. These results are decision-support measurements rather than diagnoses or predictions of healing. The primary implementation runs on a local compatible NVIDIA GPU, while a cloud GPU may be provisioned manually for selected experiments. Any commercialization claim therefore remains conditional on further validation, privacy and security controls, regulatory review, clinical workflow studies, and a sustainable deployment model.",
    "Removed promotional, diagnostic, access, and automatic-cloud claims that the current prototype and evidence do not support.",
)
replace(
    "Chapter 3 – Business Model",
    "\tThe Business Model of the study focuses on a B2B health-tech platform developed to assist endocrinologists and healthcare institutions in monitoring diabetic foot ulcers objectively. It leverages cloud-based 3DGS rendering and tissue segmentation models trained on specialized datasets (Dhar et al., 2024), targeting the aim of making clinical screening quicker, precise, and accessible.",
    "The proposed business model considers a future B2B health-technology platform for healthcare institutions, subject to successful technical and clinical validation. The prototype combines GPU-based 3DGS reconstruction with tissue characterization trained on DFUTissue (Dhar et al., 2024). Its potential value lies in structured wound documentation and review; claims of clinical precision, screening effectiveness, or accessibility require evidence beyond the current Thesis 1 prototype.",
    "Changed commercialization language from established performance to a conditional future proposition.",
)
replace(
    "Chapter 3 – Business Model",
    "To make this platform a reality, the research relies on collaborations with organizations such as the Philippine Society of Endocrinology and Metabolism, the Department of Health, and academic health-tech institutions. The main tasks include optimizing the 3DGS pipeline, managing the serialized cloud-compute queue to control GPU expenses, and conducting clinical validation. Users will have access to continuous support via frequent model updates and a self-service web interface for clinicians. Financially, the platform is supported by institutional SaaS subscriptions, pay-per-scan API billing for EMR integration, and R&D grants.",
    "A future deployment would require documented partnerships, clinical validation, secure hosting, privacy governance, technical support, and regulatory review. The present application uses an in-process sequential queue and does not provide a managed cloud-compute service, automatic cloud offloading, EMR integration, or an established subscription and billing operation. These items should be treated as roadmap assumptions to be tested rather than current resources or commitments.",
    "Removed unverified partnerships, cloud-queue operation, support, EMR billing, and financing claims.",
)

# Reference correction: remove the unmatched fabricated-looking citation and add the foundational COLMAP source.
bad_ref = find_paragraph("Wang, Y., Li, X., Zhang, H., & Chen, J. (2025). Robust incremental structure-from-motion for medical imaging applications. Journal of Medical Imaging, 12(3), 034501. https://doi.org/10.1117/1.JMI.12.3.034501")
set_paragraph_text(bad_ref, "Schönberger, J. L., & Frahm, J.-M. (2016). Structure-from-Motion Revisited. Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition, 4104–4113. https://doi.org/10.1109/CVPR.2016.445")
log("References", "Wang et al. (2025), unmatched to the in-text Liu et al. citation", "Schönberger and Frahm (2016), COLMAP Structure-from-Motion Revisited", "Made the in-text COLMAP discussion and reference list consistent.")

# Core table typo corrections.
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            txt = cell.text.strip()
            replacements = {
                "Camera poses, sparse Initialization. And Gaussian-splat representation": "Camera poses, sparse initialization, and a Gaussian-splat representation",
                "Avoids measuring directly from a smoothed mesh and supports the calculation of area, volume,d epth, width, and height": "Avoids measuring directly from a smoothed mesh and supports calculation of planar opening area, volume, depth, width, and height",
                "Required no additional surface construction": "Requires no additional surface construction",
                "Main contain insufficient point density for stable wound measurement": "May contain insufficient point density for stable wound measurement",
                "designed primarily for rendering and does not directly define a validated measurement surface": "Designed primarily for rendering and does not directly define a validated measurement surface",
            }
            if txt in replacements:
                set_cell_text(cell, replacements[txt])
                log("Chapter 3 – Trade-off tables", txt, replacements[txt], "Corrected grammar, spelling, or metric terminology.")

# Preserve the source's compatibility setting and update field codes when opened.
settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

doc.save(OUT)

# Create a separate, exhaustive change log as the second Google Docs-compatible file.
log_doc = Document()
sec = log_doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.65)
sec.right_margin = Inches(0.65)

title = log_doc.add_paragraph()
title.style = log_doc.styles["Title"]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Wound-Splat Chapters 1–3 Revision Change Log")

p = log_doc.add_paragraph()
p.add_run("Editing rule: ").bold = True
p.add_run("Only content that was incorrect, contradictory, incomplete, unsupported, or clearly malformed was changed. Unrelated wording, figures, styles, and document structure were retained.")

p = log_doc.add_paragraph()
p.add_run("Major addition: ").bold = True
p.add_run("A missing Algorithm Implementation section was added from the actual Wound-Splat source code and current prototype behavior. Longitudinal registration and Wagner/University of Texas classification support remain explicitly scheduled for Thesis 2.")

table = log_doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["No.", "Location", "Change", "Reason"]
for i, h in enumerate(headers):
    c = table.rows[0].cells[i]
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.paragraphs[0].add_run(h)
    r.bold = True

for i, ch in enumerate(changes, start=1):
    cells = table.add_row().cells
    cells[0].text = str(i)
    cells[1].text = ch["location"]
    if ch["old"] == "[Missing]":
        summary = f"Added: {ch['new']}"
    else:
        old = ch["old"].replace("\n", " ")
        new = ch["new"].replace("\n", " ")
        if len(old) > 260:
            old = old[:257] + "..."
        if len(new) > 340:
            new = new[:337] + "..."
        summary = f"Before: {old}\nAfter: {new}"
    cells[2].text = summary
    cells[3].text = ch["reason"]
    for c in cells:
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        for para in c.paragraphs:
            for run in para.runs:
                run.font.size = Pt(8)

log_doc.add_paragraph(f"Total logged changes: {len(changes)}")
log_doc.save(LOG_OUT)

print(f"saved {OUT}")
print(f"saved {LOG_OUT}")
print(f"changes {len(changes)}")
